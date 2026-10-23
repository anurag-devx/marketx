"""Documents library: reading, splitting, storing, hybrid search, replace/remove, enable/disable, answers."""
import hashlib, re
import numpy as np, pytest
for _m in ("lancedb", "llama_index.core", "trafilatura", "pypdf", "docx"):
    pytest.importorskip(_m)
from core import library as L, llm


class HashEmbed:
    """Deterministic bag-of-words embedding (no model needed)."""
    dim = 256

    def _v(self, t):
        v = np.zeros(self.dim, dtype="float32")
        for w in re.findall(r"[a-z0-9]+", t.lower()):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % self.dim] += 1
        n = np.linalg.norm(v)
        return list(v / n if n else v)

    def get_text_embedding_batch(self, texts, show_progress=False):
        return [self._v(t) for t in texts]

    def get_query_embedding(self, q):
        return self._v(q)


@pytest.fixture
def lib(tmp_path, monkeypatch):
    monkeypatch.setattr(L, "DIR", tmp_path / "library")
    monkeypatch.setattr(L, "FILES", tmp_path / "library" / "files")
    monkeypatch.setattr(L, "DB", tmp_path / "library" / "lancedb")
    monkeypatch.setattr(L, "REG", tmp_path / "library" / "sources.parquet")
    monkeypatch.setattr(L, "OPTS", tmp_path / "library" / "options.json")
    monkeypatch.setattr(L, "_tag_symbols", lambda texts: [""] * len(texts))
    L.set_embedder(HashEmbed())
    yield L
    L.set_embedder(None)


def _docx(path, heading, body):
    import docx
    d = docx.Document()
    d.add_heading(heading, 1)
    d.add_paragraph(body)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Metric", "FY26", "Order book", "Rs 1.9 lakh crore"
    d.save(path)


def test_add_search_replace_remove(lib, tmp_path):
    a = tmp_path / "hal.txt"
    a.write_text("HAL concall. Management guided order inflow of Rs 50,000 crore for FY27 and revenue growth of 15%. " * 3)
    b = tmp_path / "cement.md"
    b.write_text("UltraTech plans cement capacity of 200 MTPA by 2027. Cement prices in north India rose. " * 3)
    c = tmp_path / "note.docx"
    _docx(c, "Order book", "Bharat Electronics order book details.")
    for p in (a, b, c):
        sid, n, note = lib.add_file(p)
        assert n >= 1, note
    hits = lib.search("cement capacity expansion")
    assert hits.iloc[0].title == "cement.md"
    assert lib.search("order book lakh crore").title.iloc[0] == "note.docx"
    sid_b = lib.source_id("file", str(b))
    assert lib.add_file(b)[2] == "unchanged"                      # same content -> skipped
    b.write_text("Completely new text about steel exports. " * 5)
    lib.add_file(b)
    assert not any("UltraTech" in t for t in lib.search("UltraTech cement capacity").text)   # old passages replaced
    assert "steel" in lib.search("steel exports").text.iloc[0]
    lib.set_enabled(sid_b, False)
    assert sid_b not in set(lib.search("steel exports").source_id)
    lib.remove(sid_b)
    assert sid_b not in set(lib.registry().source_id)
    assert lib.stats()["sources"] == 2


def test_unreadable_pdf_is_flagged(lib, tmp_path):
    from pypdf import PdfWriter
    p = tmp_path / "scan.pdf"
    w = PdfWriter(); w.add_blank_page(200, 200); w.write(str(p))
    sid, n, note = lib.add_file(p)
    assert n == 0 and "scanned" in note
    assert lib.registry().set_index("source_id").at[sid, "status"] == "failed"


def test_answer_uses_docs_model_with_passages(lib, tmp_path, monkeypatch):
    p = tmp_path / "hal.txt"
    p.write_text("HAL management guided revenue growth of 15% for FY27. " * 4)
    lib.add_file(p)
    seen = {}

    def fake_chat(feature, msgs, system=None, max_tokens=0, temperature=None):
        seen.update(feature=feature, prompt=msgs[-1]["content"], system=system)
        return "Revenue growth guidance is 15% [1]."
    monkeypatch.setattr(llm, "chat", fake_chat)
    text, hits = lib.answer("What revenue growth did HAL guide?")
    assert seen["feature"] == "docs" and "[1]" in seen["prompt"] and "ONLY" in seen["system"]
    assert text.endswith("[1].") and len(hits) >= 1


def test_page_extraction_and_links():
    html = """<html><head><title>Q2 results</title></head><body><nav>Menu Home About</nav>
    <article><h1>Infosys Q2 results</h1><p>Infosys reported revenue of Rs 40,000 crore, up 8% from a year earlier.
    The company raised its full-year guidance to 3-4% growth in constant currency.</p>
    <p>Deal wins stood at $3 billion. Attrition fell to 13%.</p></article>
    <a href="/markets/a.html">A</a><a href="https://other.com/x">X</a><a href="/img.png">I</a><a href="/markets/a.html#c">A2</a></body></html>"""
    t, txt, _ = L.extract_page(html, "https://site.com/markets/q2.html")
    assert "revenue of Rs 40,000 crore" in txt and "Menu Home" not in txt
    assert L.page_links(html, "https://site.com/markets/q2.html") == ["https://site.com/markets/a.html"]


def test_online_models_blocked_by_default(monkeypatch):
    from core import config as C
    monkeypatch.setattr(C, "AI", {"features": {"ask": {"provider": "anthropic", "model": "x"}}}, raising=False)
    with pytest.raises(llm.ModelError):
        llm.chat("ask", [{"role": "user", "content": "hi"}])
    ok, msg = llm.status("ask")
    assert not ok and "switched off" in msg
    with pytest.raises(llm.ModelError):
        llm._local_url("https://evil.example.com")
    assert "sk-abc123…" in llm._redact("bad key sk-abc123XYZXYZXYZ")
