"""Documents library (RAG): your files, web pages, whole sites and news articles, searchable by the AI models.

Ingestion (LlamaIndex): read -> split into ~500-token passages (SentenceSplitter) -> embed with a local model
  (nomic-embed-text via Ollama). Each passage keeps its source, title, page/section, date and the stocks it names.
Storage + search (LanceDB, on disk in data/library/): hybrid search = meaning (vector) + keywords (full-text),
  merged by reciprocal-rank fusion. Scales to millions of passages without loading them into memory.
Answering: the best passages go to the model chosen for 'docs' (local by default), which must cite them [1], [2]...
  and say when the answer isn't in the documents.

Sources: files (PDF with selectable text, .docx, .doc via macOS textutil, .txt, .md, .html), folders, a web page,
a whole site (same-site links + sitemap, robots.txt obeyed, polite pace, page limit), and news articles
(core/articles.py). Re-adding a source replaces its passages; unchanged content is skipped."""
import hashlib, json, re, subprocess
from pathlib import Path
from urllib.parse import urljoin, urlparse, urldefrag
import numpy as np, pandas as pd
from .config import DATA
from . import atomic

DIR = DATA / "library"
FILES = DIR / "files"
DB = DIR / "lancedb"
REG = DIR / "sources.parquet"
OPTS = DIR / "options.json"
TABLE = "chunks"
EXT = {".pdf", ".docx", ".doc", ".txt", ".md", ".html", ".htm", ".rtf"}
REG_COLS = ["source_id", "kind", "title", "location", "added", "updated", "chunks", "chars", "pages", "hash", "status", "note", "enabled", "published"]
DEFAULT_OPTS = {"enabled": True, "site_page_limit": 50, "top_k": 8, "chunk_size": 512, "chunk_overlap": 64}


class LibraryError(RuntimeError):
    pass


def options():
    try:
        return {**DEFAULT_OPTS, **json.loads(OPTS.read_text())}
    except Exception:
        return dict(DEFAULT_OPTS)


def save_options(o):
    DIR.mkdir(parents=True, exist_ok=True)
    OPTS.write_text(json.dumps({**options(), **o}, indent=1))


# ---------------- registry ----------------
def registry():
    if REG.exists():
        d = pd.read_parquet(REG)
        for c in REG_COLS:
            if c not in d:
                d[c] = None
        return d[REG_COLS]
    return pd.DataFrame(columns=REG_COLS)


def _save_registry(d):
    DIR.mkdir(parents=True, exist_ok=True)
    atomic.to_parquet(d, REG)


def source_id(kind, location):
    return hashlib.sha1(f"{kind}|{location}".encode()).hexdigest()[:16]


# ---------------- reading ----------------
def read_file(path):
    """-> list of (text, page_or_section). Raises LibraryError for unreadable files."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        try:
            r = PdfReader(str(p))
        except Exception as e:
            raise LibraryError(f"PDF can't be opened ({type(e).__name__})")
        if r.is_encrypted:
            try:
                r.decrypt("")
            except Exception:
                raise LibraryError("PDF is password-protected")
        pages = []
        for i, pg in enumerate(r.pages, 1):
            try:
                t = pg.extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                pages.append((t, f"p. {i}"))
        chars = sum(len(t) for t, _ in pages)
        if chars < 50 * max(1, len(r.pages)) and chars < 2000:
            raise LibraryError("PDF has no selectable text (scanned image); not indexed")
        return pages
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        out, sec, buf = [], "start", []
        for para in d.paragraphs:
            if para.style is not None and str(para.style.name).lower().startswith("heading") and para.text.strip():
                if buf:
                    out.append(("\n".join(buf), sec)); buf = []
                sec = para.text.strip()[:80]
            if para.text.strip():
                buf.append(para.text)
        for t in d.tables:
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
            buf.append("\n".join(rows))
        if buf:
            out.append(("\n".join(buf), sec))
        return out
    if ext in (".doc", ".rtf"):
        try:
            txt = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(p)], capture_output=True, timeout=120).stdout.decode("utf-8", "ignore")
        except FileNotFoundError:
            raise LibraryError(f"{ext} files need macOS's textutil (convert to .docx or PDF on other systems)")
        if not txt.strip():
            raise LibraryError(f"{ext} file has no readable text")
        return [(txt, "")]
    if ext in (".html", ".htm"):
        import trafilatura
        txt = trafilatura.extract(p.read_text(errors="ignore"), include_tables=True) or ""
        return [(txt, "")] if txt.strip() else []
    if ext in (".txt", ".md"):
        return [(p.read_text(errors="ignore"), "")]
    raise LibraryError(f"{ext} files are not supported")


def extract_page(html, url):
    """(title, text, published) from a web page's HTML; main content only."""
    import trafilatura
    txt = trafilatura.extract(html, url=url, include_tables=True, include_comments=False, favor_precision=True) or ""
    meta = trafilatura.extract_metadata(html, default_url=url)
    title = (meta.title if meta and meta.title else "") or url
    date = meta.date if meta and meta.date else None
    return title.strip(), txt.strip(), date


def page_links(html, url):
    """Same-site links on a page (no fragments, no files we can't read)."""
    from lxml import html as LH
    try:
        tree = LH.fromstring(html)
    except Exception:
        return []
    host = urlparse(url).netloc
    out = []
    for href in tree.xpath("//a/@href"):
        u = urldefrag(urljoin(url, href))[0]
        pu = urlparse(u)
        if pu.scheme in ("http", "https") and pu.netloc == host and not re.search(r"\.(jpg|jpeg|png|gif|svg|zip|mp4|mp3|css|js|xlsx?|pptx?)$", pu.path, re.I):
            out.append(u)
    return list(dict.fromkeys(out))


# ---------------- splitting + embedding ----------------
def _splitter():
    from llama_index.core.node_parser import SentenceSplitter
    o = options()
    return SentenceSplitter(chunk_size=int(o["chunk_size"]), chunk_overlap=int(o["chunk_overlap"]))


_EMB = None


def set_embedder(e):
    """Tests inject a fake embedder here."""
    global _EMB
    _EMB = e


def embedder():
    global _EMB
    if _EMB is None:
        from . import llm
        _EMB = llm.embedder()
    return _EMB


def _tag_symbols(texts):
    try:
        from . import news
        tagger, _ = news.load_tagger()
        return [",".join(tagger.find(t[:3000])[:8]) for t in texts]
    except Exception:
        return [""] * len(texts)


def _chunks(parts, meta):
    from llama_index.core import Document
    sp = _splitter()
    rows = []
    for text, where in parts:
        text = re.sub(r"[ \t]+", " ", text or "").strip()
        if len(text) < 40:
            continue
        for n in sp.get_nodes_from_documents([Document(text=text)]):
            rows.append(dict(text=n.get_content(), where=where))
    for i, r in enumerate(rows):
        r.update(meta)
        r["id"] = f"{meta['source_id']}-{i:05d}"
    return rows


# ---------------- LanceDB ----------------
def _db():
    import lancedb
    DB.mkdir(parents=True, exist_ok=True)
    return lancedb.connect(str(DB))


def _table(db=None, create_with=None):
    db = db or _db()
    try:
        names = db.list_tables()
        names = list(getattr(names, "tables", names))
    except Exception:
        names = db.table_names()
    if TABLE in names:
        return db.open_table(TABLE)
    if create_with is None:
        return None
    return db.create_table(TABLE, data=create_with)


def _fts(t):
    if t is None:
        return
    from .filelock import locked
    with locked("library"):
        _fts_unlocked(t)


def _fts_unlocked(t):
    try:
        from lancedb.index import FTS
        t.create_index("text", config=FTS(), replace=True)
    except Exception:
        try:
            t.create_fts_index("text", replace=True)
        except Exception:
            pass


def _delete_source(t, sid):
    if t is not None and re.fullmatch(r"[0-9a-f]{16}", sid):
        from .filelock import locked
        with locked("library"):
            t.delete(f"source_id = '{sid}'")


def _store(rows):
    if not rows:
        return 0
    from .filelock import locked
    with locked("library"):
        return _store_unlocked(rows)


def _store_unlocked(rows):
    vecs = embedder().get_text_embedding_batch([r["text"] for r in rows], show_progress=False)
    sym = _tag_symbols([r["text"] for r in rows])
    data = [dict(id=r["id"], source_id=r["source_id"], kind=r["kind"], title=r["title"][:300], location=r["location"][:1000],
                 where=str(r.get("where") or "")[:80], published=str(r.get("published") or "")[:25], symbols=s,
                 text=r["text"], vector=np.asarray(v, dtype="float32")) for r, v, s in zip(rows, vecs, sym)]
    db = _db()
    t = _table(db)
    if t is None:
        t = _table(db, create_with=data)
    else:
        t.add(data)
    return len(data)


def _register(sid, kind, title, location, n, chars, pages, h, status="ok", note="", published=None):
    reg = registry()
    now = str(pd.Timestamp.now().floor("s"))
    old = reg[reg.source_id == sid]
    row = dict(source_id=sid, kind=kind, title=title[:300], location=location, added=old.added.iat[0] if len(old) else now, updated=now,
               chunks=int(n), chars=int(chars), pages=int(pages), hash=h, status=status, note=note[:300],
               enabled=bool(old.enabled.iat[0]) if len(old) and old.enabled.iat[0] is not None else True, published=str(published or ""))
    reg = pd.concat([reg[reg.source_id != sid], pd.DataFrame([row])], ignore_index=True)
    _save_registry(reg)


def add_parts(kind, location, title, parts, published=None, rebuild_fts=True):
    """Index already-extracted text. Returns (source_id, passages, skipped_unchanged)."""
    sid = source_id(kind, location)
    joined = "\n".join(t for t, _ in parts)
    h = hashlib.sha1(joined.encode("utf-8", "ignore")).hexdigest()[:16]
    reg = registry()
    old = reg[reg.source_id == sid]
    if len(old) and old.hash.iat[0] == h and old.status.iat[0] == "ok":
        return sid, int(old.chunks.iat[0]), True
    if not joined.strip():
        _register(sid, kind, title, location, 0, 0, len(parts), h, "empty", "no readable text")
        return sid, 0, False
    rows = _chunks(parts, dict(source_id=sid, kind=kind, title=title, location=location, published=published))
    t = _table()
    _delete_source(t, sid)
    n = _store(rows)
    if rebuild_fts:
        _fts(_table())
    _register(sid, kind, title, location, n, len(joined), len(parts), h, published=published)
    return sid, n, False


def add_file(path, title=None, kind="file", rebuild_fts=True):
    p = Path(path)
    try:
        parts = read_file(p)
    except LibraryError as e:
        sid = source_id(kind, str(p))
        _register(sid, kind, title or p.name, str(p), 0, 0, 0, "", "failed", str(e))
        return sid, 0, str(e)
    sid, n, skipped = add_parts(kind, str(p), title or p.name, parts, rebuild_fts=rebuild_fts)
    return sid, n, "unchanged" if skipped else ""


def add_folder(folder, log=print):
    folder = Path(folder).expanduser()
    if not folder.is_dir():
        raise LibraryError(f"Not a folder: {folder}")
    files = [p for p in sorted(folder.rglob("*")) if p.suffix.lower() in EXT and p.is_file() and not p.name.startswith(".")]
    log(f"{len(files)} readable files in {folder}")
    total = 0
    for i, p in enumerate(files, 1):
        sid, n, note = add_file(p, kind="file", rebuild_fts=False)
        total += n
        log(f"  [{i}/{len(files)}] {p.name}: {n} passages {note}")
    _fts(_table())
    return total


def add_url(url, polite=None, kind="web", rebuild_fts=True, title=None, published=None):
    from .news_sources import Polite
    polite = polite or Polite()
    html = polite.get(url).decode("utf-8", "ignore")
    t, txt, date = extract_page(html, url)
    sid, n, skipped = add_parts(kind, url, title or t, [(txt, "")], published=published or date, rebuild_fts=rebuild_fts)
    return sid, n, txt, html


def crawl_site(start, limit=None, same_path=True, log=print, polite=None):
    """Breadth-first crawl of one site from `start` (plus its sitemap), up to `limit` pages. Obeys robots.txt."""
    from .news_sources import Polite, parse_sitemap
    limit = int(limit or options()["site_page_limit"])
    polite = polite or Polite()
    base = urlparse(start)
    prefix = base.path.rsplit("/", 1)[0] + "/" if same_path and base.path not in ("", "/") else "/"
    queue, seen, done, total = [start], {start}, 0, 0
    try:                                             # sitemap pages under the same path, if the site lists them
        for sm in polite.sitemaps(f"{base.scheme}://{base.netloc}")[:3]:
            kind, items = parse_sitemap(polite.get(sm))
            locs = [x["loc"] for x in items] if kind == "urlset" else []
            for u in locs:
                if urlparse(u).path.startswith(prefix) and u not in seen and len(seen) < limit * 3:
                    seen.add(u); queue.append(u)
    except Exception:
        pass
    while queue and done < limit:
        u = queue.pop(0)
        if not polite.allowed(u):
            continue
        try:
            sid, n, txt, html = add_url(u, polite, kind="site", rebuild_fts=False)
        except Exception as e:
            log(f"  skip {u} ({type(e).__name__})"); continue
        done += 1; total += n
        log(f"  [{done}/{limit}] {u[:100]}: {n} passages")
        for l in page_links(html, u):
            if l not in seen and urlparse(l).path.startswith(prefix):
                seen.add(l); queue.append(l)
    _fts(_table())
    return done, total


def remove(sid):
    _delete_source(_table(), sid)
    reg = registry()
    row = reg[reg.source_id == sid]
    if len(row) and row.kind.iat[0] == "file":
        p = Path(row.location.iat[0])
        if FILES in p.parents:                       # only files the app copied in; never your own folders
            p.unlink(missing_ok=True)
    _save_registry(reg[reg.source_id != sid])
    _fts(_table())


def set_enabled(sid, on):
    reg = registry()
    reg.loc[reg.source_id == sid, "enabled"] = bool(on)
    _save_registry(reg)


def stats():
    reg = registry()
    t = _table()
    return dict(sources=len(reg), passages=int(t.count_rows()) if t is not None else 0,
                by_kind=reg.groupby("kind").size().to_dict() if len(reg) else {})


# ---------------- search ----------------
def search(query, k=None, kinds=None, source_ids=None, symbol=None):
    """Hybrid search. Returns a DataFrame of passages with rank (best first)."""
    t = _table()
    if t is None:
        return pd.DataFrame()
    k = int(k or options()["top_k"])
    reg = registry()
    off = set(reg[reg.enabled == False].source_id)  # noqa: E712
    where = []
    if kinds:
        where.append("kind IN (" + ",".join(f"'{x}'" for x in kinds if re.fullmatch(r"[a-z]+", x)) + ")")
    if source_ids:
        where.append("source_id IN (" + ",".join(f"'{x}'" for x in source_ids if re.fullmatch(r"[0-9a-f]{16}", x)) + ")")
    if symbol and re.fullmatch(r"[A-Z0-9&\-]+", symbol):
        where.append(f"symbols LIKE '%{symbol}%'")
    w = " AND ".join(where) or None
    cols = ["id", "source_id", "kind", "title", "location", "where", "published", "symbols", "text"]
    qv = embedder().get_query_embedding(query)
    try:
        s = t.search(np.asarray(qv, dtype="float32"), vector_column_name="vector").limit(k * 4)
        vec = (s.where(w, prefilter=True) if w else s).to_pandas()
    except Exception:
        vec = pd.DataFrame(columns=cols)
    try:
        s = t.search(query, query_type="fts").limit(k * 4)
        kw = (s.where(w, prefilter=True) if w else s).to_pandas()
    except Exception:
        kw = pd.DataFrame(columns=cols)
    score = {}
    for df in (vec, kw):
        for rank, cid in enumerate(df.id if len(df) else []):
            score[cid] = score.get(cid, 0) + 1 / (60 + rank)
    both = pd.concat([vec[[c for c in cols if c in vec]], kw[[c for c in cols if c in kw]]]).drop_duplicates("id")
    if both.empty:
        return both
    both["score"] = both.id.map(score)
    both["in_vector"] = both.id.isin(set(vec.id) if len(vec) else set())
    both["in_keywords"] = both.id.isin(set(kw.id) if len(kw) else set())
    both = both[~both.source_id.isin(off)]
    return both.sort_values("score", ascending=False).head(k).reset_index(drop=True)


DOC_SYSTEM = """You answer questions using ONLY the numbered passages from the user's own documents.
- Cite every fact with the passage number in square brackets, like [2]. Use only passages that support the statement.
- If the passages don't contain the answer, say "I couldn't find this in your documents." and say what they do cover.
- Quote figures exactly as written in the passages. Don't add outside knowledge, don't speculate, no investment advice.
- Be concise and specific: short paragraphs or a short list."""


def passages_text(hits, max_chars=1800, prefix=""):
    out = []
    for i, r in enumerate(hits.itertuples(), 1):
        loc = f"{r.title}" + (f", {r.where}" if r.where else "") + (f", {r.published[:10]}" if r.published else "")
        out.append(f"[{prefix}{i}] ({loc})\n{r.text[:max_chars]}")
    return "\n\n".join(out)


def answer(question, history=None, k=None, kinds=None, source_ids=None, symbol=None):
    """(answer text, passages used). Uses the model set for 'docs'."""
    from . import llm
    hits = search(question, k=k, kinds=kinds, source_ids=source_ids, symbol=symbol)
    if hits.empty:
        return "There's nothing in your documents library yet that matches. Add documents on the Documents page.", hits
    msgs = [m for m in (history or [])[-4:] if m["role"] in ("user", "assistant")]
    msgs.append({"role": "user", "content": f"PASSAGES:\n{passages_text(hits)}\n\nQUESTION: {question}"})
    return llm.chat("docs", msgs, system=DOC_SYSTEM, max_tokens=1200), hits


def summarise(sid, log=print, group=6):
    """Whole-source summary: summarise passages in groups, then combine (map-reduce)."""
    from . import llm
    t = _table()
    if t is None or not re.fullmatch(r"[0-9a-f]{16}", sid):
        return ""
    rows = t.search().where(f"source_id = '{sid}'").limit(100000).to_pandas().sort_values("id")
    texts = list(rows.text)
    if not texts:
        return ""
    partial = []
    for i in range(0, len(texts), group):
        block = "\n\n".join(texts[i:i + group])
        log(f"  summarising part {i // group + 1} of {(len(texts) + group - 1) // group}")
        partial.append(llm.chat("docs", [{"role": "user", "content": "Summarise the key facts, figures, guidance and risks in this part of a document "
                                                                      "in 5-8 bullet points. Keep numbers exact.\n\n" + block}], max_tokens=500))
    if len(partial) == 1:
        return partial[0]
    return llm.chat("docs", [{"role": "user", "content": "Combine these part summaries of one document into one structured summary "
                                                          "(overview, key numbers, guidance/outlook, risks). Keep numbers exact; don't add anything.\n\n"
                                                          + "\n\n".join(partial)}], max_tokens=900)


def save_upload(name, data):
    """Saves a dropped file into data/library/files/ and returns its path."""
    FILES.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w.\- ]", "_", Path(name).name)[:150]
    p = FILES / safe
    p.write_bytes(data)
    return p


# ---------------- background queue (the dashboard adds work here; the 'library-ingest' job does it) ----------------
QUEUE = DIR / "queue.json"
SUMMARIES = DIR / "summaries"


def enqueue(kind, value, **kw):
    """kind: file | folder | url | site | summary."""
    if kind not in ("file", "folder", "url", "site", "summary"):
        raise LibraryError(f"unknown task {kind}")
    from .filelock import locked
    DIR.mkdir(parents=True, exist_ok=True)
    with locked("library-queue"):
        q = json.loads(QUEUE.read_text()) if QUEUE.exists() else []
        q.append(dict(kind=kind, value=str(value), added=str(pd.Timestamp.now().floor("s")), **kw))
        atomic.write_text(QUEUE, json.dumps(q, indent=1))


def queued():
    try:
        return json.loads(QUEUE.read_text()) if QUEUE.exists() else []
    except Exception:
        return []


def run_queue(log=print):
    done = 0
    while True:
        q = queued()
        if not q:
            break
        task = q[0]
        k, v = task["kind"], task["value"]
        log(f"{k}: {v}")
        try:
            if k == "file":
                sid, n, note = add_file(v)
                log(f"  {n} passages {note}")
            elif k == "folder":
                add_folder(v, log)
            elif k == "url":
                sid, n, _, _ = add_url(v)
                log(f"  {n} passages")
            elif k == "site":
                pages, n = crawl_site(v, task.get("limit"), task.get("same_path", True), log)
                log(f"  {pages} pages, {n} passages")
            elif k == "summary":
                SUMMARIES.mkdir(parents=True, exist_ok=True)
                txt = summarise(v, log)
                (SUMMARIES / f"{v}.md").write_text(txt)
                log("  summary saved")
        except Exception as e:
            log(f"  failed: {type(e).__name__}: {str(e)[:200]}")
        from .filelock import locked
        with locked("library-queue"):
            q = queued()
            atomic.write_text(QUEUE, json.dumps(q[1:], indent=1))
        done += 1
    log(f"Library: {done} task(s) done. {stats()}")


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == "ingest":
        run_queue()
    else:
        print(stats())
